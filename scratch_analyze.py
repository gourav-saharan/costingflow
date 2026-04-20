import openpyxl
import os
import json
from datetime import datetime

file_path = "Costing from April '26.xlsx"

def analyze():
    if not os.path.exists(file_path):
        print(f"File not found: {file_path}")
        return

    wb = openpyxl.load_workbook(file_path, data_only=True)
    
    summary = {}
    
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        headers = [str(cell.value).strip() if cell.value else f"Col_{i}" for i, cell in enumerate(ws[1], 1)]
        
        real_rows = 0
        status_counts = {}
        year_counts = {}
        
        for row in ws.iter_rows(min_row=2, values_only=True):
            if any(row):
                real_rows += 1
                
                status = str(row[14]) if row[14] is not None else "None"
                status_counts[status] = status_counts.get(status, 0) + 1
                
                # Month is col 2 (index 1)
                month_val = row[1]
                year = "Unknown"
                if isinstance(month_val, datetime):
                    year = month_val.year
                elif isinstance(month_val, str):
                    try:
                        year = datetime.strptime(month_val, "%Y-%m-%d %H:%M:%S").year
                    except:
                        pass
                
                year_counts[str(year)] = year_counts.get(str(year), 0) + 1
        
        summary[sheet_name] = {
            "actual_data_rows": real_rows,
            "status_distribution": status_counts,
            "year_distribution": year_counts
        }

    print(json.dumps(summary, indent=2))

if __name__ == "__main__":
    analyze()
