import openpyxl
import os

file_path = os.path.join(os.path.dirname(__file__), "Costing from April '26.xlsx")
print(f"Loading: {file_path}")
wb = openpyxl.load_workbook(file_path, data_only=True)
print('Sheet names:', wb.sheetnames)

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    print(f'\n=== Sheet: {sheet_name} ===')
    print(f'Max row: {ws.max_row}, Max col: {ws.max_column}')
    print('First 10 rows:')
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True)):
        non_none = [str(c) if c is not None else '' for c in row]
        # trim trailing empty
        while non_none and non_none[-1] == '':
            non_none.pop()
        if any(non_none):
            print(f'  Row {i+1}: {non_none}')
