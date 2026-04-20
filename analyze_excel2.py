import openpyxl
import os
import json

file_path = os.path.join(os.path.dirname(__file__), "Costing from April '26.xlsx")
wb = openpyxl.load_workbook(file_path, data_only=True)

result = {}
for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    headers = []
    for cell in ws[1]:
        if cell.value is not None:
            headers.append(str(cell.value).strip())
        else:
            headers.append('')
    
    rows = []
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
        row_data = [str(c) if c is not None else '' for c in row]
        if any(row_data):
            rows.append(row_data)
    
    result[sheet_name] = {
        'headers': headers,
        'row_count': ws.max_row - 1,
        'col_count': ws.max_column,
        'sample_rows': rows[:5]
    }

print(json.dumps(result, indent=2, ensure_ascii=False))
